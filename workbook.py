"""Read saved worksheets and the embedded model; never refresh or modify Excel."""
from __future__ import annotations
import hashlib, io, json, os, tempfile, zipfile
from collections import deque
import pandas as pd
from openpyxl import load_workbook

SPECS=[
 ("Summary",7,2,6,()),("BOM Cost",5,2,9,()),("Sales Data",5,2,10,(0,)),
 ("Production Data",6,2,10,(0,1)),("Forecast Planning",5,2,9,(0,1,2)),
 ("Forecast Need to Order List",1,1,8,()),("Need to Order Sheet-B.O. PS",1,1,7,())]
TITLES=["Stock overview","BOM cost","Sales & transfers","Production","Forecast planning","Forecast purchasing","Order requirements"]
MAX_BYTES=25_000_000

def validate_zip(data,limit=125_000_000):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos=z.infolist()
        if len(infos)>3000 or sum(i.file_size for i in infos)>limit:
            raise ValueError("The workbook contents are too large.")
        return set(z.namelist())

def read_upload(data,filename):
    if len(data)>MAX_BYTES:
        raise ValueError("Choose a workbook or full-model package under 25 MB.")
    packaged=None
    if filename.lower().endswith(".stock.zip"):
        validate_zip(data,60_000_000)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            data=z.read("workbook.xlsx")
            packaged=json.loads(z.read("model.json"))
        if packaged.get("workbookHash")!=hashlib.sha256(data).hexdigest():
            raise ValueError("The model package does not match its workbook.")
    elif not filename.lower().endswith(".xlsx"):
        raise ValueError("Upload an .xlsx workbook or a .stock.zip package.")
    names=validate_zip(data)
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    sheets=[]
    snapshot=None
    if 'Production Data' in wb.sheetnames:
        import re
        from datetime import datetime,date
        for values in wb['Production Data'].iter_rows(min_row=1,max_row=5,values_only=True):
            for value in values:
                if isinstance(value,(datetime,date)):snapshot=value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
                elif isinstance(value,str):
                    match=re.search(r'(\d{2}[-/]\d{2}[-/]\d{4})',value)
                    if match:
                        try:snapshot=datetime.strptime(match.group(1).replace('/','-'),'%d-%m-%Y').date().isoformat()
                        except ValueError:pass

    try:
        for name,header,first,last,fill in SPECS:
            if name not in wb.sheetnames:
                raise ValueError(f"Missing worksheet: {name}")
            ws=wb[name]
            if ws.max_row>100_000:
                raise ValueError(f"Too many rows in {name}.")
            it=ws.iter_rows(min_row=header,min_col=first,max_col=last,values_only=True)
            headers=[str(v or "").replace("\n"," ").strip() for v in next(it)]
            if any(not h for h in headers):
                raise ValueError(f"Unexpected header layout in {name}.")
            rows=[];previous={}
            for raw in it:
                r=list(raw)
                if all(v is None for v in r):continue
                if name not in (SPECS[5][0],SPECS[6][0]) and any(isinstance(v,str) and (v=="Grand Total" or v.endswith(" Total")) for v in r[:4]):continue
                for i in fill:
                    if r[i] is None:r[i]=previous.get(i)
                    else:previous[i]=r[i]
                rows.append(r)
            sheets.append({"name":name,"headers":headers,"rows":rows})
    finally:
        wb.close()
    if not sheets[0]["rows"]:
        raise ValueError("Refresh and save your stock overview in Excel before uploading.")
    tables=[]
    if packaged:
        tables=packaged.get("tables",[])
        validate_model(tables)
    elif "xl/model/item.data" in names:
        from pbixray import PBIXRay
        path=None
        try:
            with tempfile.NamedTemporaryFile(suffix=".xlsx",delete=False) as tmp:
                tmp.write(data);path=tmp.name
            with PBIXRay(path) as model:
                for name in model.tables:
                    frame=model.get_table(name)
                    frame=frame[[c for c in frame.columns if c!="__XL_RowNumber"]]
                    tables.append({"name":str(name),"headers":list(frame.columns),"rows":json.loads(frame.to_json(orient="values",date_format="iso",double_precision=15))})
            validate_model(tables)
        finally:
            if path:os.unlink(path)
    return {"filename":filename,"workbookHash":hashlib.sha256(data).hexdigest(),"sheets":sheets,"tables":tables,"productionSnapshotDate":snapshot}

def validate_model(tables):
    if not isinstance(tables,list) or len(tables)>60:raise ValueError("Invalid model tables.")
    for t in tables:
        if not isinstance(t.get("name"),str) or not isinstance(t.get("headers"),list) or len(t["headers"])>200 or len(t.get("rows",[]))>100_000:
            raise ValueError("Invalid model table.")
        for r in t["rows"]:
            if not isinstance(r,list) or len(r)!=len(t["headers"]) or any(v is not None and not isinstance(v,(str,int,float)) for v in r):
                raise ValueError("Invalid model record.")

def forecast_summary(frame):
    headers=["Material","Material(text)","Total Forecast Quantity","Total Transfer Quantity","Current Stock","Pending Quantity"]
    out=[]
    for code,g in frame.groupby(frame.columns[0],sort=False,dropna=False):
        forecast=pd.to_numeric(g.iloc[:,4],errors="coerce").sum()
        transfer=pd.to_numeric(g.iloc[:,5],errors="coerce").sum()
        stocks=pd.to_numeric(g.iloc[:,6],errors="coerce").dropna()
        stock=stocks.iloc[0] if len(stocks) else 0
        out.append([code,g.iloc[0,1],float(forecast),float(transfer),float(stock),float(forecast-transfer-stock)])
    return pd.DataFrame(out,columns=headers)

def segment_map(tables):
    result={}
    def add(code,source):
        if code is None or not source:return
        name=str(source)
        if name.lower().endswith(".xlsx"):name=name[:-5]
        elif name.lower().endswith(".xls"):name=name[:-4]
        result.setdefault(str(code),set()).add(name.strip())
    for t in tables:
        if t["name"] not in ("Unique Material","BOM","BOM Cost"):continue
        h=t["headers"]
        if "Material" not in h or "Source.Name" not in h:continue
        for r in t["rows"]:add(r[h.index("Material")],r[h.index("Source.Name")])
    boms=[t for t in tables if "Material" in t["headers"] and "Component" in t["headers"]]
    for _ in range(20):
        changed=False
        for t in boms:
            mi=t["headers"].index("Material");ci=t["headers"].index("Component")
            for r in t["rows"]:
                if not r[ci]:continue
                old=result.setdefault(str(r[ci]),set());before=len(old)
                old.update(result.get(str(r[mi]),set()))
                changed|=len(old)!=before
        if not changed:break
    return result

def segment_filter(frame,mapping,segment):
    if segment=="All segments":return frame
    key=next((k for k in ["Material","Material Code","Material Number","SKU Code","Component"] if k in frame.columns),None)
    if key is None:return frame.iloc[0:0] if "Date" not in frame.columns else frame
    return frame[frame[key].map(lambda c:segment in mapping.get(str(c),set()))]

def related_tables(tables,code,entire=False):
    if entire:return tables
    if code is None:return [{**t,"rows":[]} for t in tables]
    children=set();parents=set();boms=[t for t in tables if "Material" in t["headers"] and "Component" in t["headers"]]
    for t in boms:
        mi=t["headers"].index("Material");ci=t["headers"].index("Component")
        parents.update(r[mi] for r in t["rows"] if r[ci]==code and r[mi])
    queue=deque([code]);seen=set()
    while queue:
        current=queue.popleft()
        if current in seen:continue
        seen.add(current)
        for t in boms:
            mi=t["headers"].index("Material");ci=t["headers"].index("Component")
            for r in t["rows"]:
                if r[mi]==current and r[ci] and r[ci]!=code and r[ci] not in children:
                    children.add(r[ci]);queue.append(r[ci])
    linked={code}|children|parents;out=[];dates=set()
    for t in tables:
        h=t["headers"];key=next((k for k in ["Material","Material Code","Material Number","SKU Code"] if k in h),None)
        if not key:
            out.append({**t,"rows":[]});continue
        ki=h.index(key);ci=h.index("Component") if "Component" in h else -1;rows=[]
        for r in t["rows"]:
            if r[ki] in linked or (ci>=0 and r[ci]==code):
                relation="Selected material" if r[ki]==code else "BOM component" if r[ki] in children else "Product using selected component"
                rows.append(r+[relation])
                for i,c in enumerate(h):
                    if any(s in c.lower() for s in ["date","attribute","valid from","created on","changed on","last change"]) and isinstance(r[i],str) and len(r[i])>=10:
                        dates.add(r[i][:10])
        out.append({**t,"headers":h+["Relationship to selection"],"rows":rows})
    for t in out:
        if t["name"]=="Calendar":
            original=next(x for x in tables if x["name"]=="Calendar");di=original["headers"].index("Date")
            t["rows"]=[r for r in original["rows"] if str(r[di])[:10] in dates]
    return out

PRODUCTION_REVISION=1
