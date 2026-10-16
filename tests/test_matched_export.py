import base64,io
import pytest
from openpyxl import load_workbook
from test_stock_assistant import accounts
from stock_accounts.workspace_insights import handle_request


def fixture(store,admin):
    data={'sheets':[{'name':'Summary','headers':['Material','Qty'],'rows':[['A',5]]}],
          'tables':[{'name':'Stock Data','headers':['Material','Qty','Description'],
                     'rows':[['A',5.0,'=literal'],['B',99,'unrelated']]}]}
    store.save_dataset(admin,'master.xlsx',data)
    return {'action':'matched_export','version':store.get_dataset(admin)['version'],
            'scope':'related','tab':0,'row':['A',5],'col':1,'table':'Stock Data',
            'headers':['Material','Qty'],'rows':[['A',5]]}


def test_viewer_can_download_only_value_matched_table_without_broad_permission(accounts):
    store,admin,viewer=accounts;event=fixture(store,admin)
    assert not store.identity(viewer)['permissions']['download']
    result=handle_request(store,viewer,event)
    wb=load_workbook(io.BytesIO(base64.b64decode(result['content'])))
    assert wb.sheetnames==['Matched records']
    assert list(wb.active.values)==[('Material','Qty'),('A',5)]
    assert not store.identity(viewer)['permissions']['download']
    with pytest.raises(PermissionError):handle_request(store,viewer,{'action':'export'})
    assert any(e['action']=='export' and e['resource']=='Value matched records: Stock Data' for e in store.access_log(admin))


@pytest.mark.parametrize('changes',[
 {'scope':'all'},{'exportTables':True},{'rows':[['B',99]]},
 {'row':['B',99]},{'table':'accounts'},{'version':'stale'},{'col':-1},
 {'headers':['Password']},{'rows':[['A',{'forged':'value'}]]},
])
def test_forged_and_broad_exports_are_rejected(accounts,changes):
    store,admin,viewer=accounts;event=fixture(store,admin)
    with pytest.raises(ValueError):handle_request(store,viewer,{**event,**changes})


def test_download_rechecks_revoked_session_and_preserves_literal_strings(accounts):
    store,admin,viewer=accounts;event=fixture(store,admin)
    event.update(headers=['Description'],rows=[['=literal']])
    result=handle_request(store,viewer,event)
    wb=load_workbook(io.BytesIO(base64.b64decode(result['content'])))
    assert wb.active['A2'].data_type=='s'
    store.update_viewer(admin,'viewer_user',active=False)
    with pytest.raises(PermissionError):handle_request(store,viewer,event)


def test_forecast_product_total_click_with_blank_stock_can_download_matches(accounts):
    store,admin,viewer=accounts
    data={'sheets':[{'name':'Empty','headers':['Material'],'rows':[]} for _ in range(4)]+[
        {'name':'Forecast Planning','headers':['Material','Name','Year','Month','Forecast','Transfer','Stock','Pending'],
         'rows':[['A','Product',2026,'Apr',10,2,None,8],['A','Product',2026,'May',20,3,None,17]]}],
         'tables':[{'name':'Stock Data','headers':['Material','Qty'],'rows':[['A',5],['B',99]]}]}
    store.save_dataset(admin,'forecast.xlsx',data)
    event={'action':'matched_export','version':store.get_dataset(viewer)['version'],'scope':'related',
           'tab':4,'row':['A','Product',30,5,None,25],'col':5,'table':'Stock Data','headers':['Material','Qty'],'rows':[['A',5]]}
    assert handle_request(store,viewer,event)['filename']=='Matched_records.xlsx'
