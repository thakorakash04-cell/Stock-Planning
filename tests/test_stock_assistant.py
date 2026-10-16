import tempfile
from pathlib import Path
from unittest.mock import patch
import streamlit as st
import pytest
from stock_accounts.store import Store,AccessDenied
from stock_accounts.workspace import handle_request

@pytest.fixture
def accounts():
    with tempfile.TemporaryDirectory() as folder:
        store=Store('sqlite:///'+str(Path(folder)/'db'));store.bootstrap_admin('admin_user','Test-admin-password!123');admin=store.login('admin_user','Test-admin-password!123')
        store.create_viewer(admin,'viewer_user','Test-viewer-password!123');viewer=store.login('viewer_user','Test-viewer-password!123')
        store.save_dataset(admin,'stock.xlsx',{'sheets':[{'name':'Stock','headers':['Material','Qty'],'rows':[['A',3],['B',9]]}],'tables':[]})
        with patch('streamlit.session_state',{}):yield store,admin,viewer

def test_permission_consent_revision_and_context_are_enforced(accounts):
    store,admin,viewer=accounts
    status=handle_request(store,viewer,{'action':'ai_status'})
    assert 'key' not in status and 'api_key' not in status
    event={'action':'ask_ai','version':status['version'],'question':'Why this value?','context':{'table':'Stock','column':'Qty','row':['A',3]}}
    with patch('grounded_ai.generate') as provider:
        with pytest.raises(ValueError,match='sharing'):handle_request(store,viewer,event)
        with pytest.raises(ValueError,match='changed'):handle_request(store,viewer,{**event,'consent':True,'version':'stale'})
        provider.assert_not_called()
    with patch('grounded_ai.settings',return_value=('key','model')),patch('grounded_ai.generate') as provider:
        with pytest.raises(ValueError,match='selected record'):handle_request(store,viewer,{**event,'consent':True,'context':{'table':'Stock','column':'Qty','row':['A',999]}})
        provider.assert_not_called()
    store.update_permissions(admin,'viewer_user',{'use_tools':False});viewer=store.login('viewer_user','Test-viewer-password!123')
    with patch('grounded_ai.generate') as provider:
        for action in ['ai_status','ask_ai','ai_clear']:
            with pytest.raises(AccessDenied):handle_request(store,viewer,{**event,'action':action,'consent':True})
        provider.assert_not_called()

def test_success_keeps_private_history_and_generic_audit_then_clear(accounts):
    store,admin,viewer=accounts
    status=handle_request(store,viewer,{'action':'ai_status'})
    event={'action':'ask_ai','version':status['version'],'question':'Private question','context':None,'consent':True}
    with patch('stock_accounts.assistant.ask',return_value={'answer':'Private answer','evidence':[]}) as ask:
        assert handle_request(store,viewer,event)['answer']=='Private answer'
        handle_request(store,viewer,{**event,'question':'follow up'})
        assert ask.call_args.args[2][0]['text']=='Private question'
        handle_request(store,viewer,{'action':'ai_clear','context':None})
        handle_request(store,viewer,event)
        assert ask.call_args.args[2]==[]
    logs=store.access_log(admin)
    assert any(e['action']=='tool_use' and e['resource']=='Ask AI' for e in logs)
    assert 'Private' not in str(logs)

def test_forecast_product_totals_deduct_stock_once(accounts):
    from stock_accounts.assistant import stock_tables
    store,admin,viewer=accounts
    sheets=[{'name':'Empty '+str(i),'headers':['Material'],'rows':[]} for i in range(4)]
    sheets.append({'name':'Forecast Planning','headers':['Material','Description','Month','Other','Forecast','Transfer','Stock'],'rows':[['A','Item','Jan',None,10,2,3],['A','Item','Feb',None,20,4,3]]})
    store.save_dataset(admin,'forecast.xlsx',{'sheets':sheets,'tables':[]})
    tables,version,name=stock_tables(store,viewer)
    summary=next(t for t in tables if t['name']=='Forecast Planning — product totals')
    assert summary['rows']==[['A','Item',30.0,6.0,3.0,21.0]]

def test_admin_downloads_all_matches_and_all_source_columns(accounts):
    import base64,io
    from openpyxl import load_workbook
    from grounded_ai import execute_query
    store,admin,viewer=accounts
    rows=[['=NOT_A_FORMULA',i,'X'] for i in range(1,31)]+[['Excluded',0,'Y']]
    store.save_dataset(admin,'export.xlsx',{'sheets':[{'name':'Stock','headers':['Material','Qty','Segment'],'rows':rows}],'tables':[]})
    status=handle_request(store,admin,{'action':'ai_status'})
    query={'table':'T1','filters':[{'column':'Qty','op':'gt','value':0}],'columns':['Material','Qty'],'limit':20}
    from stock_accounts.assistant import stock_tables
    tables,_,_=stock_tables(store,admin)
    result={'answer':'30 matches','evidence':[{'citation':'Q1',**execute_query(tables,query)}]}
    with patch('stock_accounts.assistant.ask',return_value=result):
        answer=handle_request(store,admin,{'action':'ask_ai','version':status['version'],'question':'Private test question','consent':True})
    evidence=answer['evidence'][0];assert evidence['shown_rows']==20
    file=handle_request(store,admin,{'action':'ai_export','export_id':evidence['export_id'],'version':status['version'],'query':{'table':'accounts'}})
    workbook=load_workbook(io.BytesIO(base64.b64decode(file['content'])))
    assert workbook['Result table'].max_row==31
    assert workbook['Matching records'].max_row==31
    assert [c.value for c in workbook['Matching records'][1]]==['Source row','Material','Qty','Segment']
    assert workbook['Matching records']['B2'].value=='=NOT_A_FORMULA'
    assert workbook['Matching records']['B2'].data_type=='s'
    assert any(e['action']=='export' and e['resource']=='Ask AI matching records' for e in store.access_log(admin))
    assert 'Private test question' not in str(store.access_log(admin))
    # A viewer with ordinary download permission must still be denied.
    store.update_permissions(admin,'viewer_user',{'download':True})
    viewer=store.login('viewer_user','Test-viewer-password!123')
    with pytest.raises(AccessDenied,match='administrators'):
        handle_request(store,viewer,{'action':'ai_export','export_id':evidence['export_id'],'version':status['version']})
    with pytest.raises(ValueError,match='expired'):
        handle_request(store,admin,{'action':'ai_export','export_id':'forged','version':status['version']})
    store.save_dataset(admin,'new.xlsx',{'sheets':[{'name':'Stock','headers':['Material'],'rows':[['new']]}],'tables':[]})
    with pytest.raises(ValueError,match='changed'):
        handle_request(store,admin,{'action':'ai_export','export_id':evidence['export_id'],'version':status['version']})

def test_grouped_download_contains_result_and_underlying_records(accounts):
    import base64,io
    from openpyxl import load_workbook
    from grounded_ai import execute_query
    from stock_accounts.assistant import stock_tables
    store,admin,viewer=accounts
    tables,version,_=stock_tables(store,admin)
    query={'table':'T1','metrics':[{'column':'Qty','op':'sum'}]}
    result={'answer':'Total 12','evidence':[{'citation':'Q1',**execute_query(tables,query)}]}
    with patch('stock_accounts.assistant.ask',return_value=result):
        answer=handle_request(store,admin,{'action':'ask_ai','question':'Total quantity','version':version,'consent':True})
    file=handle_request(store,admin,{'action':'ai_export','export_id':answer['evidence'][0]['export_id'],'version':version})
    workbook=load_workbook(io.BytesIO(base64.b64decode(file['content'])))
    assert workbook['Result table']['A2'].value==12
    assert workbook['Matching records'].max_row==3
