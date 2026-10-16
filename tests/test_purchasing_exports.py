import base64,io
from openpyxl import load_workbook
import pytest
from test_stock_assistant import accounts
from test_planning import dataset
from stock_accounts.workspace_insights import handle_request


def setup(store,admin):
    data=dataset()
    data['sheets']=[{'name':'Forecast Need to Order List','headers':['Material','Need'],'rows':[['X',10],['Y',20]]},
                    {'name':'Need to Order Sheet-B.O. PS','headers':['Material','Need'],'rows':[['X',30]]},
                    {'name':'Summary','headers':['Material','Need'],'rows':[['Secret',999]]}]
    store.save_dataset(admin,'master.xlsx',data)
    return store.get_dataset(admin)['version']


@pytest.mark.parametrize('source,rows',[('forecast_purchasing',[['Y',20]]),('order_requirements_saved',[['X',30]])])
def test_viewer_can_download_approved_purchasing_tables_only(accounts,source,rows):
    store,admin,viewer=accounts;version=setup(store,admin)
    event={'action':'planning_table_export','source':source,'version':version,'headers':['Material','Need'],'rows':rows}
    result=handle_request(store,viewer,event)
    wb=load_workbook(io.BytesIO(base64.b64decode(result['content'])))
    assert list(wb.active.values)==[('Material','Need'),tuple(rows[0])]
    assert not store.identity(viewer)['permissions']['download']
    with pytest.raises(PermissionError):handle_request(store,viewer,{'action':'export'})
    for change in [{'source':'Summary'},{'version':'old'},{'rows':[['Secret',999]]},{'headers':['Password']}]:
        with pytest.raises(ValueError):handle_request(store,viewer,{**event,**change})
    assert any(e['action']=='export' and e['resource'] in ('Forecast purchasing','Order requirements') for e in store.access_log(admin))


def test_calculated_requirements_export_is_owned_by_login_and_current_plan(accounts):
    store,admin,viewer=accounts;version=setup(store,admin)
    plan=handle_request(store,viewer,{'action':'planning_calculate','version':version,'entries':[{'code':'A','quantity':10}]})
    event={'action':'planning_table_export','source':'order_requirements_plan','version':version,'plan_id':plan['plan_id'],
           'headers':plan['headers'],'rows':plan['rows']}
    wb=load_workbook(io.BytesIO(base64.b64decode(handle_request(store,viewer,event)['content'])))
    assert list(wb.active.values)[1:]==[tuple(r) for r in plan['rows']]
    with pytest.raises(ValueError,match='expired'):handle_request(store,admin,event)
    with pytest.raises(ValueError,match='expired'):handle_request(store,viewer,{**event,'plan_id':'forged'})
    handle_request(store,viewer,{'action':'planning_calculate','version':version,'entries':[{'code':'B','quantity':10}]})
    with pytest.raises(ValueError,match='expired'):handle_request(store,viewer,event)
    store.update_viewer(admin,'viewer_user',active=False)
    with pytest.raises(PermissionError):handle_request(store,viewer,event)
