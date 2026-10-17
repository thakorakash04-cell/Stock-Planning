from unittest.mock import patch
import base64,io
import pytest
from openpyxl import load_workbook
from test_stock_assistant import accounts
from stock_accounts.custom_insight import create
from stock_accounts.workspace_insights import handle_request

TABLE={'id':'T1','name':'Stock','headers':['Material','Qty'],'rows':[['A',3],['B',9]]}
VIEW={'view':0,'title':'Stock overview','table':'Stock'}
def plan(table='T1'):
 return {'chart':{'title':'Requested quantities','type':'bar','label':'Material','value':'Qty','query':{'table':table,'columns':['Material','Qty'],'sort_by':'Qty','descending':True,'limit':10}}}

def test_language_chart_values_and_interpretation_use_validated_current_view():
 with patch('stock_accounts.custom_insight.settings',return_value=('key','model')),patch('stock_accounts.custom_insight.generate',side_effect=[plan(),{'finding':'B has the highest displayed quantity.','action':'Review source units.'}]) as provider:
  result=create([TABLE],VIEW,{'knowledge':{'rules':'PS differs from availability'}},'Show quantities',[])
 assert result['chart']['evidence']['records'][0]['Qty']==9
 assert result['chart']['finding']=='B has the highest displayed quantity.'
 assert provider.call_args_list[1].args[3]['chart']['evidence']['matched_records']==2
 with patch('stock_accounts.custom_insight.settings',return_value=('key','model')),patch('stock_accounts.custom_insight.generate',return_value=plan('UNAVAILABLE')):
  with pytest.raises(ValueError,match='wrong planning view'):create([TABLE],VIEW,{},'Show data',[])

def test_missing_fields_clarify_and_interpretation_failure_keeps_graph():
 with patch('stock_accounts.custom_insight.settings',return_value=('key','model')),patch('stock_accounts.custom_insight.generate',return_value={'chart':None,'clarification':'Which date column?'}) as provider:
  assert create([TABLE],VIEW,{},'Show a trend',[])['clarification']=='Which date column?'
  assert provider.call_count==1
 with patch('stock_accounts.custom_insight.settings',return_value=('key','model')),patch('stock_accounts.custom_insight.generate',side_effect=[plan(),ValueError('Quota unavailable')]):
  result=create([TABLE],VIEW,{},'Show quantities',[])
 assert result['chart']['evidence']['matched_records']==2 and 'unavailable' in result['chart']['data_quality']

def report(store,actor):
 version=store.get_dataset(actor)['version']
 with patch('stock_accounts.insights.build_report',return_value={'sections':[{'view':0,'title':'Stock overview','charts':[]}],'answer':'Report'}):
  result=handle_request(store,actor,{'action':'generate_insights','version':version,'consent':True})
 return {'action':'add_insight_chart','version':version,'consent':True,'report_id':result['report_id'],'view':0,'question':'Show quantities'}

def test_generated_report_chart_is_read_only_downloadable_and_duplicate_safe(accounts):
 store,admin,viewer=accounts;event=report(store,admin);before=store.get_dataset(admin)
 with patch('stock_accounts.custom_insight.settings',return_value=('key','model')),patch('stock_accounts.custom_insight.generate',side_effect=[plan(),{'finding':'Review ranking'},plan(),{'finding':'Review ranking'}]):
  result=handle_request(store,admin,event)
  duplicate=handle_request(store,admin,event)
 assert result['chart']['custom'] and duplicate.get('clarification')
 assert store.get_dataset(admin)==before
 e=result['chart']['evidence']
 records=handle_request(store,admin,{'action':'insight_records','source_id':e['source_id'],'version':event['version']})
 assert records['total']==2
 file=handle_request(store,admin,{'action':'ai_export','export_id':e['export_id'],'version':event['version']})
 assert load_workbook(io.BytesIO(base64.b64decode(file['content'])))['Matching records'].max_row==3
 assert 'Show quantities' not in str(store.access_log(admin))

def test_reports_permissions_consent_and_workbook_versions_are_enforced(accounts):
 store,admin,viewer=accounts;event=report(store,admin)
 with patch('stock_accounts.custom_insight.generate') as provider:
  for change in [{'consent':False},{'report_id':'forged'},{'view':99},{'version':'old'}]:
   with pytest.raises(ValueError):handle_request(store,admin,{**event,**change})
  with pytest.raises(ValueError):handle_request(store,viewer,event)
  provider.assert_not_called()
 store.update_permissions(admin,'viewer_user',{'menu_8':False});viewer=store.login('viewer_user','Test-viewer-password!123')
 with pytest.raises(PermissionError):handle_request(store,viewer,event)
