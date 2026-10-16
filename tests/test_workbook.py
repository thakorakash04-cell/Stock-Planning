import unittest
import pandas as pd
from workbook import forecast_summary,segment_map,segment_filter,related_tables

class WorkbookTests(unittest.TestCase):
    def test_stock_deducted_once_and_negative_balance_retained(self):
        frame=pd.DataFrame([["A","One",2026,"April",150,100,20,30],["A","One",2026,"May",200,50,20,130],["B","Two",2026,"April",10,20,2,-12]],columns=["Material","Name","Year","Month","Forecast","Transfer","Stock","Pending"])
        out=forecast_summary(frame)
        self.assertEqual(out.loc[0,"Pending Quantity"],180)
        self.assertEqual(out.loc[1,"Pending Quantity"],-12)
        self.assertEqual(out["Current Stock"].sum(),22)
    def test_source_names_components_and_related_rows(self):
        tables=[{"name":"Unique Material","headers":["Material","Source.Name"],"rows":[["A","Ostomy Care.XLSX"],["B","General Wound Care.xlsx"]]}, {"name":"BOM","headers":["Material","Component","Source.Name"],"rows":[["A","C","Ostomy Care.XLSX"],["B","C","General Wound Care.xlsx"]]}, {"name":"DP MRP","headers":["SKU Code","Segment"],"rows":[["A","Marketing category"]]}, {"name":"Stock Data","headers":["Material","Current Stock"],"rows":[["A",5],["C",20],["Z",2]]}]
        mapping=segment_map(tables)
        self.assertEqual(mapping["C"],{"Ostomy Care","General Wound Care"})
        self.assertNotIn("Marketing category",{s for v in mapping.values() for s in v})
        matched=related_tables(tables,"A")
        stock=next(t for t in matched if t["name"]=="Stock Data")
        self.assertEqual({r[0] for r in stock["rows"]},{"A","C"})
        self.assertEqual(related_tables(tables,"A",True),tables)
        bom=pd.DataFrame(tables[1]["rows"],columns=tables[1]["headers"])
        self.assertEqual(segment_filter(bom,mapping,"Ostomy Care")["Material"].tolist(),["A"])

def test_model_import_keeps_precision_needed_for_ps_whole_number_rounding():
    import io,zipfile
    from unittest.mock import patch,MagicMock
    from openpyxl import Workbook
    from workbook import read_upload,SPECS
    wb=Workbook();wb.remove(wb.active)
    for name,header,first,last,_ in SPECS:
        ws=wb.create_sheet(name)
        for col in range(first,last+1):ws.cell(header,col,'Column '+str(col))
        ws.cell(header+1,first,'A')
    out=io.BytesIO();wb.save(out)
    with zipfile.ZipFile(out,'a') as z:z.writestr('xl/model/item.data',b'mocked model')
    frame=pd.DataFrame({'Component quantity':[0.0014999999998]})
    model=MagicMock();model.__enter__.return_value=model;model.tables=['BOM'];model.get_table.return_value=frame
    with patch('pbixray.PBIXRay',return_value=model):
        value=read_upload(out.getvalue(),'master.xlsx')['tables'][0]['rows'][0][0]
    assert value==0.0014999999998
    assert round(value*1000)==1  # Rounding to 10 decimals would produce 2.
